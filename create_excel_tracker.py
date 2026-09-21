import calendar
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_exact_tracker_excel():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Typography & Styles
    FONT_FAMILY = "Arial"
    
    # Border definitions
    THIN_BORDER_GRAY = Side(border_style="thin", color="000000")
    MEDIUM_BORDER = Side(border_style="medium", color="000000")
    
    BORDER_STANDARD = Border(
        left=THIN_BORDER_GRAY, right=THIN_BORDER_GRAY,
        top=THIN_BORDER_GRAY, bottom=THIN_BORDER_GRAY
    )
    BORDER_SUNDAY = Border(
        left=MEDIUM_BORDER, right=MEDIUM_BORDER,
        top=THIN_BORDER_GRAY, bottom=THIN_BORDER_GRAY
    )
    BORDER_NOTE = Border(
        left=Side(border_style="thin", color="E2E8F0"),
        right=Side(border_style="thin", color="E2E8F0"),
        top=Side(border_style="thin", color="E2E8F0"),
        bottom=Side(border_style="thin", color="E2E8F0")
    )

    # Color definitions matching the PDF exactly
    # 5 Group colors for 15 protocol rows (3 rows each)
    GROUP_COLORS_HEX = [
        "FFEBEB",  # Soft Red/Pink: RGB(255, 235, 235)
        "EBF5FF",  # Soft Sky/Blue: RGB(235, 245, 255)
        "EBFFEB",  # Soft Mint/Green: RGB(235, 255, 235)
        "FFFFEB",  # Soft Cream/Yellow: RGB(255, 255, 235)
        "F5EBFF",  # Soft Lavender/Purple: RGB(245, 235, 255)
    ]
    
    # Sunday deeper tint in protocol rows (max(0, c-25))
    SUNDAY_GROUP_COLORS_HEX = [
        "E6D2D2",  # RGB(230, 210, 210)
        "D2DCE6",  # RGB(210, 220, 230)
        "D2E6D2",  # RGB(210, 230, 210)
        "E6E6D2",  # RGB(230, 230, 210)
        "DCD2E6",  # RGB(220, 210, 230)
    ]

    SUNDAY_HEADER_FILL = PatternFill(start_color="D2D2D2", end_color="D2D2D2", fill_type="solid") # RGB(210, 210, 210)
    NOTES_FILL = PatternFill(start_color="FFFFE6", end_color="FFFFE6", fill_type="solid") # RGB(255, 255, 230)
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Quotes from trackergenerator01.py
    MONTH_QUOTES = {
        9: "Win the morning, win the day.",
        10: "Discipline equals freedom.",
        11: "Focus on the process.",
        12: "Stay consistent."
    }

    # Only September 2026 to December 2026
    target_months = [
        (2026, 9),
        (2026, 10),
        (2026, 11),
        (2026, 12)
    ]

    weekdays_abbr = ['M', 'T', 'W', 'Th', 'F', 'Sa', 'Su']

    for year, month in target_months:
        month_name = calendar.month_name[month]
        sheet_title = f"{month_name[:3]} {year}"
        ws = wb.create_sheet(title=sheet_title)
        
        # Gridlines and Page setup for Landscape A4 print
        ws.views.sheetView[0].showGridLines = True
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1

        days_in_month = calendar.monthrange(year, month)[1]
        first_weekday = calendar.monthrange(year, month)[0] # 0 = Monday, 6 = Sunday
        month_weekdays = [weekdays_abbr[(first_weekday + d) % 7] for d in range(days_in_month)]

        # Columns:
        # Col A (1): S.No.
        # Col B (2): Protocols
        # Col C (3): Target
        # Col D.. (4 .. 3+days_in_month): Days 1..N
        # Col after last day: Success %
        first_day_col = 4
        last_day_col = 3 + days_in_month
        success_col = last_day_col + 1

        last_day_letter = get_column_letter(last_day_col)
        success_col_letter = get_column_letter(success_col)

        # Set Column Widths
        ws.column_dimensions['A'].width = 6.5
        ws.column_dimensions['B'].width = 28
        ws.column_dimensions['C'].width = 8.5
        for col_idx in range(first_day_col, last_day_col + 1):
            c_let = get_column_letter(col_idx)
            ws.column_dimensions[c_let].width = 3.6
        ws.column_dimensions[success_col_letter].width = 11.5

        # Row 1: Top spacing
        ws.row_dimensions[1].height = 8

        # --- Row 2: Title, Quote, and @surajvansh ---
        ws.row_dimensions[2].height = 24
        ws["A2"] = "Habit Tracker"
        ws["A2"].font = Font(name=FONT_FAMILY, size=18, bold=True, color="000000")
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

        # Quote centered in Row 2 across middle columns
        quote_text = MONTH_QUOTES.get(month, "Win the morning, win the day.")
        quote_start_col = 8
        quote_end_col = last_day_col - 5
        ws.merge_cells(start_row=2, start_column=quote_start_col, end_row=2, end_column=quote_end_col)
        quote_cell = ws.cell(row=2, column=quote_start_col)
        quote_cell.value = f'"{quote_text}"'
        quote_cell.font = Font(name=FONT_FAMILY, size=11, bold=True, color="000000")
        quote_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Top Right: @surajvansh
        handle_start_col = last_day_col - 4
        ws.merge_cells(start_row=2, start_column=handle_start_col, end_row=2, end_column=success_col)
        handle_cell = ws.cell(row=2, column=handle_start_col)
        handle_cell.value = "@surajvansh"
        handle_cell.font = Font(name=FONT_FAMILY, size=13, bold=True, color="000000")
        handle_cell.alignment = Alignment(horizontal="right", vertical="center")

        # --- Row 3: Month Year and Days Dashboard ---
        ws.row_dimensions[3].height = 18
        ws["A3"] = f"{month_name} {year}"
        ws["A3"].font = Font(name=FONT_FAMILY, size=12, bold=True, color="000000")
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells(start_row=3, start_column=quote_start_col, end_row=3, end_column=quote_end_col)
        dash_cell = ws.cell(row=3, column=quote_start_col)
        dash_cell.value = f"{days_in_month} Days Dashboard"
        dash_cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
        dash_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Row 4: spacing
        ws.row_dimensions[4].height = 6

        # --- Rows 5 & 6: Protocol Table Header ---
        ws.row_dimensions[5].height = 15
        ws.row_dimensions[6].height = 15

        # Merge S.No.
        ws.merge_cells("A5:A6")
        ws["A5"] = "S.No."
        ws["A5"].font = Font(name=FONT_FAMILY, size=8, bold=True)
        ws["A5"].alignment = Alignment(horizontal="center", vertical="center")
        for r in (5, 6):
            ws[f"A{r}"].border = BORDER_STANDARD

        # Merge Protocols
        ws.merge_cells("B5:B6")
        ws["B5"] = "Protocols"
        ws["B5"].font = Font(name=FONT_FAMILY, size=8, bold=True)
        ws["B5"].alignment = Alignment(horizontal="center", vertical="center")
        for r in (5, 6):
            ws[f"B{r}"].border = BORDER_STANDARD

        # Merge Target
        ws.merge_cells("C5:C6")
        ws["C5"] = "Target"
        ws["C5"].font = Font(name=FONT_FAMILY, size=8, bold=True)
        ws["C5"].alignment = Alignment(horizontal="center", vertical="center")
        for r in (5, 6):
            ws[f"C{r}"].border = BORDER_STANDARD

        # Day Number and Weekday Headers
        for d in range(1, days_in_month + 1):
            col = first_day_col + (d - 1)
            c_let = get_column_letter(col)
            is_sunday = (month_weekdays[d - 1] == 'Su')
            header_fill = SUNDAY_HEADER_FILL if is_sunday else WHITE_FILL

            # Row 5: Day number
            cell_d = ws[f"{c_let}5"]
            cell_d.value = d
            cell_d.font = Font(name=FONT_FAMILY, size=8, bold=True)
            cell_d.alignment = Alignment(horizontal="center", vertical="center")
            cell_d.fill = header_fill
            cell_d.border = BORDER_SUNDAY if is_sunday else BORDER_STANDARD

            # Row 6: Weekday abbreviation
            cell_w = ws[f"{c_let}6"]
            cell_w.value = month_weekdays[d - 1]
            cell_w.font = Font(name=FONT_FAMILY, size=7, bold=(is_sunday))
            cell_w.alignment = Alignment(horizontal="center", vertical="center")
            cell_w.fill = header_fill
            cell_w.border = BORDER_SUNDAY if is_sunday else BORDER_STANDARD

        # Merge Success %
        ws.merge_cells(f"{success_col_letter}5:{success_col_letter}6")
        ws[f"{success_col_letter}5"] = "Success %"
        ws[f"{success_col_letter}5"].font = Font(name=FONT_FAMILY, size=7.5, bold=True)
        ws[f"{success_col_letter}5"].alignment = Alignment(horizontal="center", vertical="center")
        for r in (5, 6):
            ws[f"{success_col_letter}{r}"].border = BORDER_STANDARD

        # --- Rows 7 to 21: 15 Protocol Rows (5 color groups of 3 rows) ---
        first_habit_row = 7
        last_habit_row = 21

        for i in range(1, 16):
            row = first_habit_row + (i - 1)
            ws.row_dimensions[row].height = 14.5
            group_idx = (i - 1) // 3
            group_fill_hex = GROUP_COLORS_HEX[group_idx]
            sunday_fill_hex = SUNDAY_GROUP_COLORS_HEX[group_idx]

            row_fill = PatternFill(start_color=group_fill_hex, end_color=group_fill_hex, fill_type="solid")
            row_sunday_fill = PatternFill(start_color=sunday_fill_hex, end_color=sunday_fill_hex, fill_type="solid")

            # Col A: S.No.
            cell_sn = ws[f"A{row}"]
            cell_sn.value = i
            cell_sn.font = Font(name=FONT_FAMILY, size=8)
            cell_sn.alignment = Alignment(horizontal="center", vertical="center")
            cell_sn.fill = row_fill
            cell_sn.border = BORDER_STANDARD

            # Col B: Protocols (blank, ready for habit name)
            cell_proto = ws[f"B{row}"]
            cell_proto.font = Font(name=FONT_FAMILY, size=8.5)
            cell_proto.alignment = Alignment(horizontal="left", vertical="center")
            cell_proto.fill = row_fill
            cell_proto.border = BORDER_STANDARD

            # Col C: Target
            cell_tgt = ws[f"C{row}"]
            cell_tgt.font = Font(name=FONT_FAMILY, size=8)
            cell_tgt.alignment = Alignment(horizontal="center", vertical="center")
            cell_tgt.fill = row_fill
            cell_tgt.border = BORDER_STANDARD

            # Cols D..: Day cells
            for d in range(1, days_in_month + 1):
                col = first_day_col + (d - 1)
                c_let = get_column_letter(col)
                is_sunday = (month_weekdays[d - 1] == 'Su')
                cell_day = ws[f"{c_let}{row}"]
                cell_day.font = Font(name=FONT_FAMILY, size=8)
                cell_day.alignment = Alignment(horizontal="center", vertical="center")
                cell_day.fill = row_sunday_fill if is_sunday else row_fill
                cell_day.border = BORDER_SUNDAY if is_sunday else BORDER_STANDARD

            # Col Success %: Dynamic formula supporting numbers (1) and checks (✓, x)
            cell_succ = ws[f"{success_col_letter}{row}"]
            first_day_let = get_column_letter(first_day_col)
            last_d_let = get_column_letter(last_day_col)
            
            # Excel formula: calculates completion rate if protocol has entries
            calc_expr = f'(COUNT({first_day_let}{row}:{last_d_let}{row}) + COUNTIF({first_day_let}{row}:{last_d_let}{row}, "✓") + COUNTIF({first_day_let}{row}:{last_d_let}{row}, "x") + COUNTIF({first_day_let}{row}:{last_d_let}{row}, "X"))'
            cell_succ.value = f'=IF(ISBLANK(B{row}), "", IF(ISNUMBER(C{row}), IF(C{row}>0, {calc_expr}/C{row}, 0), {calc_expr}/{days_in_month}))'
            cell_succ.font = Font(name=FONT_FAMILY, size=8)
            cell_succ.alignment = Alignment(horizontal="center", vertical="center")
            cell_succ.fill = row_fill
            cell_succ.border = BORDER_STANDARD
            cell_succ.number_format = "0.0%"

        # --- Row 22: DAILY TOTAL SCORE ---
        score_row = 22
        ws.row_dimensions[score_row].height = 20
        
        # Merge A22:C22
        ws.merge_cells(f"A{score_row}:C{score_row}")
        score_label = ws[f"A{score_row}"]
        score_label.value = "DAILY TOTAL SCORE"
        score_label.font = Font(name=FONT_FAMILY, size=8, bold=True)
        score_label.alignment = Alignment(horizontal="right", vertical="center")
        for c in ("A", "B", "C"):
            ws[f"{c}{score_row}"].border = BORDER_STANDARD

        for d in range(1, days_in_month + 1):
            col = first_day_col + (d - 1)
            c_let = get_column_letter(col)
            cell_score = ws[f"{c_let}{score_row}"]
            # Sum up completed habits for the day
            cell_score.value = f'=IF((COUNT({c_let}{first_habit_row}:{c_let}{last_habit_row}) + COUNTIF({c_let}{first_habit_row}:{c_let}{last_habit_row}, "✓") + COUNTIF({c_let}{first_habit_row}:{c_let}{last_habit_row}, "x") + COUNTIF({c_let}{first_habit_row}:{c_let}{last_habit_row}, "X")) > 0, SUM({c_let}{first_habit_row}:{c_let}{last_habit_row}) + COUNTIF({c_let}{first_habit_row}:{c_let}{last_habit_row}, "✓") + COUNTIF({c_let}{first_habit_row}:{c_let}{last_habit_row}, "x") + COUNTIF({c_let}{first_habit_row}:{c_let}{last_habit_row}, "X"), "")'
            cell_score.font = Font(name=FONT_FAMILY, size=8, bold=True)
            cell_score.alignment = Alignment(horizontal="center", vertical="center")
            cell_score.border = BORDER_STANDARD

        # Success cell for daily score row
        succ_summary = ws[f"{success_col_letter}{score_row}"]
        succ_summary.value = f'=IF(COUNT({success_col_letter}{first_habit_row}:{success_col_letter}{last_habit_row}) > 0, AVERAGE({success_col_letter}{first_habit_row}:{success_col_letter}{last_habit_row}), "")'
        succ_summary.font = Font(name=FONT_FAMILY, size=8, bold=True)
        succ_summary.alignment = Alignment(horizontal="center", vertical="center")
        succ_summary.border = BORDER_STANDARD
        succ_summary.number_format = "0.0%"

        # --- Row 23: Spacing ---
        ws.row_dimensions[23].height = 4

        # --- Row 24 & 25: Stakes / Reward & Punishment ---
        ws.row_dimensions[24].height = 14
        ws.row_dimensions[25].height = 14

        ws.merge_cells(f"A24:{success_col_letter}24")
        punish_cell = ws["A24"]
        punish_cell.value = "IF SUCCESS < 80%, PUNISHMENT: ________________________________________________"
        punish_cell.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color="C80000") # RGB(200, 0, 0)
        punish_cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells(f"A25:{success_col_letter}25")
        reward_cell = ws["A25"]
        reward_cell.value = "IF SUCCESS > 90%, REWARD: __________________________________________________"
        reward_cell.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color="007800") # RGB(0, 120, 0)
        reward_cell.alignment = Alignment(horizontal="left", vertical="center")

        # --- Row 26: Spacing ---
        ws.row_dimensions[26].height = 4

        # --- Sleep Tracking Header (Rows 27 & 28) ---
        ws.row_dimensions[27].height = 14
        ws.row_dimensions[28].height = 14

        # Sleep Tracking Label merged across A27:C28
        ws.merge_cells("A27:C28")
        sleep_hdr = ws["A27"]
        sleep_hdr.value = "Sleep Tracking"
        sleep_hdr.font = Font(name=FONT_FAMILY, size=8.5, bold=True)
        sleep_hdr.alignment = Alignment(horizontal="center", vertical="center")
        for r in (27, 28):
            for c in ("A", "B", "C"):
                ws[f"{c}{r}"].border = BORDER_STANDARD

        for d in range(1, days_in_month + 1):
            col = first_day_col + (d - 1)
            c_let = get_column_letter(col)
            is_sunday = (month_weekdays[d - 1] == 'Su')
            header_fill = SUNDAY_HEADER_FILL if is_sunday else WHITE_FILL

            # Row 27: Day number
            cell_d = ws[f"{c_let}27"]
            cell_d.value = d
            cell_d.font = Font(name=FONT_FAMILY, size=8, bold=True)
            cell_d.alignment = Alignment(horizontal="center", vertical="center")
            cell_d.fill = header_fill
            cell_d.border = BORDER_SUNDAY if is_sunday else BORDER_STANDARD

            # Row 28: Weekday
            cell_w = ws[f"{c_let}28"]
            cell_w.value = month_weekdays[d - 1]
            cell_w.font = Font(name=FONT_FAMILY, size=7, bold=is_sunday)
            cell_w.alignment = Alignment(horizontal="center", vertical="center")
            cell_w.fill = header_fill
            cell_w.border = BORDER_SUNDAY if is_sunday else BORDER_STANDARD

        # Sleep Rows (Rows 29 to 33: 9hrs, 8hrs, 7hrs, 6hrs, 5hrs)
        sleep_labels = ["9hrs", "8hrs", "7hrs", "6hrs", "5hrs"]
        for idx, s_label in enumerate(sleep_labels):
            s_row = 29 + idx
            ws.row_dimensions[s_row].height = 13.5

            # Merge A..C for label
            ws.merge_cells(f"A{s_row}:C{s_row}")
            lbl_cell = ws[f"A{s_row}"]
            lbl_cell.value = s_label
            lbl_cell.font = Font(name=FONT_FAMILY, size=8)
            lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
            for c in ("A", "B", "C"):
                ws[f"{c}{s_row}"].border = BORDER_STANDARD

            for d in range(1, days_in_month + 1):
                col = first_day_col + (d - 1)
                c_let = get_column_letter(col)
                is_sunday = (month_weekdays[d - 1] == 'Su')
                cell_s = ws[f"{c_let}{s_row}"]
                cell_s.font = Font(name=FONT_FAMILY, size=8)
                cell_s.alignment = Alignment(horizontal="center", vertical="center")
                if is_sunday:
                    cell_s.fill = SUNDAY_HEADER_FILL
                    cell_s.border = BORDER_SUNDAY
                else:
                    cell_s.border = BORDER_STANDARD

        # --- Row 34: Spacing ---
        ws.row_dimensions[34].height = 4

        # --- Row 35: Notes Section Label ---
        ws.row_dimensions[35].height = 15
        ws["A35"] = "NOTES:"
        ws["A35"].font = Font(name=FONT_FAMILY, size=8.5, bold=True)
        ws["A35"].alignment = Alignment(horizontal="left", vertical="center")

        # Rows 36 to 38: Notes Yellow Box
        for r_note in range(36, 39):
            ws.row_dimensions[r_note].height = 13
        ws.merge_cells(f"A36:{last_day_letter}38")
        note_box = ws["A36"]
        note_box.fill = NOTES_FILL
        note_box.font = Font(name=FONT_FAMILY, size=8.5)
        note_box.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        for r in range(36, 39):
            for c in range(1, last_day_col + 1):
                ws.cell(row=r, column=c).fill = NOTES_FILL
                ws.cell(row=r, column=c).border = BORDER_NOTE

        # --- Row 39: Spacing ---
        ws.row_dimensions[39].height = 6

        # --- Row 40: Monthly Targets Label ---
        ws.row_dimensions[40].height = 16
        ws["A40"] = "MONTHLY TARGETS"
        ws["A40"].font = Font(name=FONT_FAMILY, size=8.5, bold=True)
        ws["A40"].alignment = Alignment(horizontal="left", vertical="center")

        # Rows 41 to 44: Monthly Targets (2 columns: 1-4 and 5-8)
        col2_target_start = 16 # Start around column P
        for i in range(1, 5):
            t_row = 40 + i
            ws.row_dimensions[t_row].height = 13

            # Column 1 (Targets 1-4)
            ws.merge_cells(start_row=t_row, start_column=1, end_row=t_row, end_column=col2_target_start - 2)
            t1_cell = ws.cell(row=t_row, column=1)
            t1_cell.value = f"{i}. ____________________ : ____________________"
            t1_cell.font = Font(name=FONT_FAMILY, size=7.5)
            t1_cell.alignment = Alignment(horizontal="left", vertical="center")

            # Column 2 (Targets 5-8)
            i2 = i + 4
            ws.merge_cells(start_row=t_row, start_column=col2_target_start, end_row=t_row, end_column=last_day_col)
            t2_cell = ws.cell(row=t_row, column=col2_target_start)
            t2_cell.value = f"{i2}. ____________________ : ____________________"
            t2_cell.font = Font(name=FONT_FAMILY, size=7.5)
            t2_cell.alignment = Alignment(horizontal="left", vertical="center")

    output_filename = "Habit_Tracker_Sep_Dec_2026.xlsx"
    wb.save(output_filename)
    # Also save as Habit_Tracker_2026_2027.xlsx to replace the old multi-year file with the clean requested 4-month tracker
    wb.save("Habit_Tracker_2026_2027.xlsx")
    print(f"Workbook successfully saved to {output_filename} and Habit_Tracker_2026_2027.xlsx")

if __name__ == "__main__":
    generate_exact_tracker_excel()
